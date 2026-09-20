import re
from pathlib import Path
from types import SimpleNamespace
from typing import Any, List, Tuple

from app.validation.enums import DocumentType
from app.validation.models import ValidationResult
from .base import BaseParser
from .exceptions import DocumentParsingError
from .models import ParsedDocument


# Standard section titles commonly found in documents, reports, and resumes
COMMON_HEADINGS = {
    "education",
    "projects",
    "technical skills",
    "skills",
    "experience",
    "work experience",
    "professional experience",
    "achievements",
    "certifications",
    "achievements & certifications",
    "position of responsibility",
    "positions of responsibility",
    "extracurricular activities",
    "interests",
    "summary",
    "professional summary",
    "overview",
    "introduction",
    "conclusion",
    "abstract",
    "methodology",
    "results",
    "discussion",
    "background",
    "table of contents",
    "references",
}


class LightweightDocumentItem:
    """
    Emulates Docling's DocumentElement/Item interface without importing
    heavy deep-learning libraries into memory.
    """

    def __init__(
        self,
        text: str,
        label: str = "text",
        level: int = 1,
        page_no: int = 1,
        table_markdown: str | None = None,
    ):
        self.text = text.strip()
        self.label = label
        self.level = level
        self.prov = [SimpleNamespace(page_no=page_no)]
        self._table_markdown = table_markdown

    def export_to_markdown(self, **kwargs) -> str:
        if self._table_markdown:
            return self._table_markdown
        return self.text

    def export_to_dataframe(self):
        return None


class LightweightRawDocument:
    """
    Lightweight document container compatible with SectionExtractor.
    """

    def __init__(
        self,
        items: List[Tuple[LightweightDocumentItem, int]],
        pages_count: int,
        markdown: str = "",
        text: str = "",
    ):
        self._items = items
        self.pages = list(range(1, max(1, pages_count) + 1))
        self._markdown = markdown
        self._text = text

    def iterate_items(self):
        for item, level in self._items:
            yield item, level

    def export_to_markdown(self) -> str:
        return self._markdown

    def export_to_text(self) -> str:
        return self._text


class LightweightParser(BaseParser):
    """
    Zero-vision, low-memory parser designed specifically for <= 512 MB deployments.
    Supports PDF, DOCX, PPTX, XLSX, HTML, TXT, and Markdown with < 30 MB peak RAM.
    """

    def parse(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        try:
            doc_type = validation.detection.document_type

            if doc_type == DocumentType.PDF:
                return self._parse_pdf(file_path, validation)
            elif doc_type == DocumentType.DOCX:
                return self._parse_docx(file_path, validation)
            elif doc_type == DocumentType.PPTX:
                return self._parse_pptx(file_path, validation)
            elif doc_type == DocumentType.XLSX:
                return self._parse_xlsx(file_path, validation)
            elif doc_type == DocumentType.CSV:
                return self._parse_csv(file_path, validation)
            elif doc_type == DocumentType.HTML:
                return self._parse_html(file_path, validation)
            else:  # TXT, MARKDOWN, UNKNOWN fallback
                return self._parse_text(file_path, validation)

        except Exception as exc:
            raise DocumentParsingError(
                f"Lightweight parser failed to parse document: {exc}"
            ) from exc

    # =========================================================
    # PDF Parser (using pypdfium2, ~15-25 MB RAM)
    # =========================================================

    def _parse_pdf(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(file_path)
        items: List[Tuple[LightweightDocumentItem, int]] = []
        full_text_parts: List[str] = []
        full_md_parts: List[str] = []

        current_major_section = None

        try:
            pages_count = len(pdf)

            for page_idx in range(pages_count):
                page_no = page_idx + 1
                page = pdf[page_idx]
                tp = page.get_textpage()
                try:
                    # Extract text objects
                    text_objs = []
                    for obj in page.get_objects():
                        if isinstance(obj, pdfium._helpers.pageobjects.PdfTextObj):
                            try:
                                obj.textpage = tp
                                t = obj.extract().strip()
                                if t:
                                    text_objs.append({
                                        "text": t,
                                        "size": round(obj.get_font_size(), 1),
                                        "bounds": obj.get_bounds(),
                                    })
                            except Exception:
                                pass

                    if not text_objs:
                        # Fallback to get_text_range if no text objects
                        raw_text = tp.get_text_range().strip()
                        if raw_text:
                            item = LightweightDocumentItem(
                                text=raw_text,
                                label="text",
                                level=1,
                                page_no=page_no,
                            )
                            items.append((item, 1))
                            full_text_parts.append(raw_text)
                            full_md_parts.append(raw_text)
                        continue

                    # Determine body font baseline (most frequent font size)
                    sizes = [o["size"] for o in text_objs if len(o["text"]) > 2]
                    body_size = max(set(sizes), key=sizes.count) if sizes else 10.0

                    # Sort objects top-to-bottom, left-to-right
                    text_objs.sort(key=lambda o: (-o["bounds"][3], o["bounds"][0]))

                    # Group objects on roughly the same horizontal line
                    lines = []
                    current_line = []
                    current_top = None

                    for obj in text_objs:
                        top = obj["bounds"][3]
                        if current_top is None or abs(top - current_top) < 3.5:
                            current_line.append(obj)
                            if current_top is None:
                                current_top = top
                        else:
                            lines.append(current_line)
                            current_line = [obj]
                            current_top = top
                    if current_line:
                        lines.append(current_line)

                    for line_objs in lines:
                        line_objs.sort(key=lambda o: o["bounds"][0])
                        line_text = " ".join(o["text"] for o in line_objs).strip()
                        if not line_text:
                            continue

                        max_font = max(o["size"] for o in line_objs)
                        clean_lower = line_text.lower().strip()

                        # Structural heading heuristics
                        is_known_heading = clean_lower in COMMON_HEADINGS
                        is_title = max_font >= (body_size + 4.0) or (page_no == 1 and max_font >= 16.0)
                        is_section = is_known_heading or (max_font >= (body_size + 1.0) and len(line_text) < 60)
                        is_numbered_section = bool(re.match(r"^\d+(\.\d+)*\s+[A-Z]", line_text)) and len(line_text) < 70

                        # Detect subheadings under sections like Projects
                        is_subheading = False
                        if current_major_section in ("projects", "work experience", "experience"):
                            if (
                                len(line_text) < 50
                                and not line_text.startswith("http")
                                and not line_text.startswith("•")
                                and not line_text.startswith("-")
                                and not any(k in clean_lower for k in ("github", "project link", "live link", "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"))
                                and (max_font >= body_size or max_font >= 10.8)
                                and not is_section
                            ):
                                is_subheading = True

                        if is_title or is_section or is_numbered_section:
                            level = 1
                            if is_known_heading:
                                current_major_section = clean_lower

                            item = LightweightDocumentItem(
                                text=line_text,
                                label="section_header",
                                level=level,
                                page_no=page_no,
                            )
                            items.append((item, level))
                            full_md_parts.append(f"\n# {line_text}\n")

                        elif is_subheading:
                            level = 2
                            item = LightweightDocumentItem(
                                text=line_text,
                                label="section_header",
                                level=level,
                                page_no=page_no,
                            )
                            items.append((item, level))
                            full_md_parts.append(f"\n## {line_text}\n")

                        else:
                            # Check for table-like line
                            is_table = "|" in line_text or (len(line_objs) >= 3 and all(len(o["text"]) < 30 for o in line_objs))
                            label = "table" if is_table else "text"
                            item = LightweightDocumentItem(
                                text=line_text,
                                label=label,
                                level=1,
                                page_no=page_no,
                                table_markdown=line_text if is_table else None,
                            )
                            items.append((item, 1))
                            full_md_parts.append(line_text)

                        full_text_parts.append(line_text)
                finally:
                    try:
                        tp.close()
                    except Exception:
                        pass
                    try:
                        page.close()
                    except Exception:
                        pass
        finally:
            pdf.close()

        raw_doc = LightweightRawDocument(
            items=items,
            pages_count=pages_count,
            markdown="\n".join(full_md_parts),
            text="\n".join(full_text_parts),
        )

        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=pages_count,
        )

    # =========================================================
    # DOCX Parser (using python-docx, ~5 MB RAM)
    # =========================================================

    def _parse_docx(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        import docx

        doc = docx.Document(file_path)
        items: List[Tuple[LightweightDocumentItem, int]] = []
        md_parts: List[str] = []
        text_parts: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = (p.style.name or "").lower()
            if "heading 1" in style_name or "title" in style_name:
                item = LightweightDocumentItem(text=text, label="section_header", level=1, page_no=1)
                items.append((item, 1))
                md_parts.append(f"\n# {text}\n")
            elif "heading 2" in style_name:
                item = LightweightDocumentItem(text=text, label="section_header", level=2, page_no=1)
                items.append((item, 2))
                md_parts.append(f"\n## {text}\n")
            elif "heading" in style_name:
                item = LightweightDocumentItem(text=text, label="section_header", level=3, page_no=1)
                items.append((item, 3))
                md_parts.append(f"\n### {text}\n")
            else:
                item = LightweightDocumentItem(text=text, label="text", level=1, page_no=1)
                items.append((item, 1))
                md_parts.append(text)
            text_parts.append(text)

        # Tables
        for table in doc.tables:
            rows_data = []
            for row in table.rows:
                rows_data.append([c.text.strip().replace("\n", " ") for c in row.cells])
            if rows_data:
                header = rows_data[0]
                md_table = "| " + " | ".join(header) + " |\n"
                md_table += "| " + " | ".join(["---"] * len(header)) + " |\n"
                for r in rows_data[1:]:
                    md_table += "| " + " | ".join(r) + " |\n"

                item = LightweightDocumentItem(
                    text=md_table,
                    label="table",
                    level=1,
                    page_no=1,
                    table_markdown=md_table,
                )
                items.append((item, 1))
                md_parts.append("\n" + md_table + "\n")
                text_parts.append(md_table)

        raw_doc = LightweightRawDocument(
            items=items,
            pages_count=1,
            markdown="\n".join(md_parts),
            text="\n".join(text_parts),
        )

        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=1,
        )

    # =========================================================
    # PPTX Parser (using python-pptx, ~5 MB RAM)
    # =========================================================

    def _parse_pptx(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        from pptx import Presentation

        prs = Presentation(file_path)
        items: List[Tuple[LightweightDocumentItem, int]] = []
        md_parts: List[str] = []
        text_parts: List[str] = []
        pages_count = len(prs.slides)

        for idx, slide in enumerate(prs.slides):
            page_no = idx + 1
            title = ""
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()
                item = LightweightDocumentItem(
                    text=title,
                    label="section_header",
                    level=1,
                    page_no=page_no,
                )
                items.append((item, 1))
                md_parts.append(f"\n# Slide {page_no}: {title}\n")
                text_parts.append(title)
            else:
                item = LightweightDocumentItem(
                    text=f"Slide {page_no}",
                    label="section_header",
                    level=1,
                    page_no=page_no,
                )
                items.append((item, 1))

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        txt = p.text.strip()
                        if txt and txt != title:
                            item = LightweightDocumentItem(
                                text=txt,
                                label="text",
                                level=1,
                                page_no=page_no,
                            )
                            items.append((item, 1))
                            md_parts.append(txt)
                            text_parts.append(txt)
                elif shape.has_table:
                    t = shape.table
                    rows_data = []
                    for row in t.rows:
                        rows_data.append([c.text.strip() for c in row.cells])
                    if rows_data:
                        header = rows_data[0]
                        md_table = "| " + " | ".join(header) + " |\n"
                        md_table += "| " + " | ".join(["---"] * len(header)) + " |\n"
                        for r in rows_data[1:]:
                            md_table += "| " + " | ".join(r) + " |\n"
                        item = LightweightDocumentItem(
                            text=md_table,
                            label="table",
                            level=1,
                            page_no=page_no,
                            table_markdown=md_table,
                        )
                        items.append((item, 1))
                        md_parts.append("\n" + md_table + "\n")
                        text_parts.append(md_table)

        raw_doc = LightweightRawDocument(
            items=items,
            pages_count=pages_count,
            markdown="\n".join(md_parts),
            text="\n".join(text_parts),
        )

        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=pages_count,
        )

    # =========================================================
    # XLSX Parser (using openpyxl, ~10 MB RAM)
    # =========================================================

    def _parse_xlsx(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        items: List[Tuple[LightweightDocumentItem, int]] = []
        md_parts: List[str] = []
        text_parts: List[str] = []
        sheet_names = wb.sheetnames

        for idx, name in enumerate(sheet_names):
            page_no = idx + 1
            ws = wb[name]

            # Sheet header
            item = LightweightDocumentItem(
                text=name,
                label="section_header",
                level=1,
                page_no=page_no,
            )
            items.append((item, 1))
            md_parts.append(f"\n# Sheet: {name}\n")
            text_parts.append(f"Sheet: {name}")

            rows = list(ws.iter_rows(values_only=True))
            if rows:
                header = [str(c or "").strip() for c in rows[0]]
                md_table = "| " + " | ".join(header) + " |\n"
                md_table += "| " + " | ".join(["---"] * len(header)) + " |\n"
                for r in rows[1:]:
                    vals = [str(c or "").strip() for c in r]
                    if any(vals):
                        md_table += "| " + " | ".join(vals) + " |\n"

                item = LightweightDocumentItem(
                    text=md_table,
                    label="table",
                    level=1,
                    page_no=page_no,
                    table_markdown=md_table,
                )
                items.append((item, 1))
                md_parts.append(md_table)
                text_parts.append(md_table)

        wb.close()

        raw_doc = LightweightRawDocument(
            items=items,
            pages_count=len(sheet_names),
            markdown="\n".join(md_parts),
            text="\n".join(text_parts),
        )

        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=len(sheet_names),
        )

    # =========================================================
    # CSV Parser (standard library, negligible memory)
    # =========================================================

    def _parse_csv(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        import csv

        with file_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as csv_file:
            rows = list(csv.reader(csv_file))

        if not rows:
            raise ValueError("CSV document is empty.")

        header = [cell.strip() for cell in rows[0]]
        markdown_rows = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(["---"] * len(header)) + " |",
        ]
        for row in rows[1:]:
            values = [cell.strip() for cell in row]
            if any(values):
                markdown_rows.append("| " + " | ".join(values) + " |")
        markdown = "\n".join(markdown_rows)
        title = file_path.stem.replace("_", " ").replace("-", " ").title()
        section = LightweightDocumentItem(text=title, label="section_header", level=1, page_no=1)
        table = LightweightDocumentItem(
            text=markdown,
            label="table",
            level=1,
            page_no=1,
            table_markdown=markdown,
        )
        raw_doc = LightweightRawDocument(
            items=[(section, 1), (table, 1)],
            pages_count=1,
            markdown=f"# {title}\n\n{markdown}",
            text=f"{title}\n{markdown}",
        )
        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=1,
        )

    # =========================================================
    # HTML Parser (using beautifulsoup4, ~2 MB RAM)
    # =========================================================

    def _parse_html(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        from bs4 import BeautifulSoup

        content = file_path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(content, "html.parser")
        items: List[Tuple[LightweightDocumentItem, int]] = []
        md_parts: List[str] = []
        text_parts: List[str] = []

        for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table"]):
            t = tag.get_text(strip=True)
            if not t:
                continue

            if tag.name == "h1":
                item = LightweightDocumentItem(text=t, label="section_header", level=1, page_no=1)
                items.append((item, 1))
                md_parts.append(f"\n# {t}\n")
            elif tag.name == "h2":
                item = LightweightDocumentItem(text=t, label="section_header", level=2, page_no=1)
                items.append((item, 2))
                md_parts.append(f"\n## {t}\n")
            elif tag.name in ("h3", "h4"):
                item = LightweightDocumentItem(text=t, label="section_header", level=3, page_no=1)
                items.append((item, 3))
                md_parts.append(f"\n### {t}\n")
            elif tag.name == "table":
                rows = []
                for tr in tag.find_all("tr"):
                    cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
                    if cells:
                        rows.append(cells)
                if rows:
                    header = rows[0]
                    md_table = "| " + " | ".join(header) + " |\n"
                    md_table += "| " + " | ".join(["---"] * len(header)) + " |\n"
                    for r in rows[1:]:
                        md_table += "| " + " | ".join(r) + " |\n"
                    item = LightweightDocumentItem(
                        text=md_table,
                        label="table",
                        level=1,
                        page_no=1,
                        table_markdown=md_table,
                    )
                    items.append((item, 1))
                    md_parts.append("\n" + md_table + "\n")
            else:
                item = LightweightDocumentItem(text=t, label="text", level=1, page_no=1)
                items.append((item, 1))
                md_parts.append(t)
            text_parts.append(t)

        raw_doc = LightweightRawDocument(
            items=items,
            pages_count=1,
            markdown="\n".join(md_parts),
            text="\n".join(text_parts),
        )

        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=1,
        )

    # =========================================================
    # Markdown & TXT Parser (~1 MB RAM)
    # =========================================================

    def _parse_text(
        self,
        file_path: Path,
        validation: ValidationResult,
    ) -> ParsedDocument:
        content = file_path.read_text(encoding="utf-8", errors="replace")
        lines = content.splitlines()
        items: List[Tuple[LightweightDocumentItem, int]] = []
        md_parts: List[str] = []
        text_parts: List[str] = []

        has_initial_header = any(re.match(r"^#{1,3}\s+", l.strip()) for l in lines[:5])
        if not has_initial_header:
            title = file_path.stem.replace("_", " ").replace("-", " ").title()
            item = LightweightDocumentItem(text=title, label="section_header", level=1, page_no=1)
            items.append((item, 1))
            md_parts.append(f"# {title}\n")

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            h1_match = re.match(r"^#\s+(.+)$", line_str)
            h2_match = re.match(r"^##\s+(.+)$", line_str)
            h3_match = re.match(r"^###\s+(.+)$", line_str)

            if h1_match:
                title = h1_match.group(1).strip()
                item = LightweightDocumentItem(text=title, label="section_header", level=1, page_no=1)
                items.append((item, 1))
                md_parts.append(f"\n# {title}\n")
            elif h2_match:
                title = h2_match.group(1).strip()
                item = LightweightDocumentItem(text=title, label="section_header", level=2, page_no=1)
                items.append((item, 2))
                md_parts.append(f"\n## {title}\n")
            elif h3_match:
                title = h3_match.group(1).strip()
                item = LightweightDocumentItem(text=title, label="section_header", level=3, page_no=1)
                items.append((item, 3))
                md_parts.append(f"\n### {title}\n")
            else:
                item = LightweightDocumentItem(text=line_str, label="text", level=1, page_no=1)
                items.append((item, 1))
                md_parts.append(line_str)
            text_parts.append(line_str)

        raw_doc = LightweightRawDocument(
            items=items,
            pages_count=1,
            markdown="\n".join(md_parts),
            text="\n".join(text_parts),
        )

        return ParsedDocument(
            source=file_path,
            validation=validation,
            markdown=raw_doc.export_to_markdown(),
            text=raw_doc.export_to_text(),
            metadata={},
            raw_document=raw_doc,
            pages=1,
        )
